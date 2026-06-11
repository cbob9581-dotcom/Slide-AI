"""
PPT/PPTX 文档解析处理
使用 LibreOffice 转 PDF + python-pptx 提取备注
"""
import os
import subprocess


def parse_ppt(path: str, out_dir: str) -> dict:
    """解析 PPT/PPTX 文件"""
    # 1. 提取演讲者备注
    notes_map = _extract_notes(path)

    # 2. LibreOffice 转 PDF
    try:
        pdf_path = _convert_to_pdf(path, out_dir)
    except Exception as e:
        return {"status": "error", "message": f"PPT to PDF conversion failed: {e}"}

    # 3. 复用 pdf_handler 解析 PDF
    from pdf_handler import parse_pdf
    result = parse_pdf(pdf_path, out_dir)

    if result.get("status") != "ok":
        return result

    # 4. 注入备注
    for page in result.get("pages", []):
        n = page["page_number"]
        page["speaker_notes"] = notes_map.get(n, "")

    return result


def _convert_to_pdf(path: str, out_dir: str) -> str:
    """使用 LibreOffice 将 PPT 转为 PDF"""
    # 尝试多个可能的 LibreOffice 路径
    libreoffice_paths = [
        "libreoffice",
        "soffice",
        "/usr/bin/libreoffice",
        "/usr/bin/soffice",
        "C:\\Program Files\\LibreOffice\\program\\soffice.exe",
        "C:\\Program Files (x86)\\LibreOffice\\program\\soffice.exe",
    ]

    libreoffice = None
    for path_candidate in libreoffice_paths:
        try:
            subprocess.run(
                [path_candidate, "--version"],
                capture_output=True,
                timeout=5,
                check=False,
            )
            libreoffice = path_candidate
            break
        except (FileNotFoundError, subprocess.TimeoutExpired):
            continue

    if not libreoffice:
        raise RuntimeError(
            "LibreOffice not found. Please install LibreOffice to process PPT files.\n"
            "Download: https://www.libreoffice.org/download/"
        )

    subprocess.run(
        [
            libreoffice,
            "--headless",
            "--convert-to", "pdf",
            path,
            "--outdir", out_dir,
        ],
        check=True,
        timeout=120,
        capture_output=True,
    )

    base = os.path.splitext(os.path.basename(path))[0]
    pdf_path = os.path.join(out_dir, base + ".pdf")

    if not os.path.exists(pdf_path):
        raise RuntimeError(f"PDF conversion output not found: {pdf_path}")

    return pdf_path


def _extract_notes(path: str) -> dict:
    """使用 python-pptx 提取演讲者备注"""
    try:
        from pptx import Presentation
        prs = Presentation(path)
        notes = {}
        for i, slide in enumerate(prs.slides, 1):
            try:
                if slide.has_notes_slide:
                    text = slide.notes_slide.notes_text_frame.text
                    if text.strip():
                        notes[i] = text.strip()
            except Exception:
                continue
        return notes
    except ImportError:
        return {}  # python-pptx 不可用时跳过备注
    except Exception:
        return {}
